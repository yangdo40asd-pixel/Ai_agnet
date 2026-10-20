"""선택 종목 분석 결과(DashboardResult/AnalysisResult)를 편집 가능한 XLSX 바이트로 변환하는 모듈.

- 외부 API를 다시 호출하지 않고 화면과 동일한 AnalysisResult 객체만 사용한다.
- 6자리 종목코드의 선행 0은 문자열(write_string)로 보존한다.
- 결측값(None)은 0으로 바꾸지 않고 빈 셀(write_blank)과 상태/사유 열로 구분한다.
- 수식 주입(Formula Injection) 방지를 위해 strings_to_formulas=False 및 위험 접두 문자 이스케이프를 적용한다.
- XlsxWriter 네이티브 차트를 실제 수치 셀 범위에 연결하여, 엑셀에서 셀 값 수정 시 차트가 자동 갱신되도록 구성한다.
"""

from __future__ import annotations

import io
import math
import re
import xml.etree.ElementTree as ET
import zipfile
from decimal import Decimal

import xlsxwriter

from src.data.news import safe_news_url
from src.models import AnalysisResult, MetricValue
from src.ui.formatting import (
    format_eok,
    format_multiple,
    format_per_share_krw,
    format_percent,
    format_price,
    format_timestamp,
)

SHEET_NAMES: tuple[str, ...] = (
    "요약",
    "핵심지표",
    "일별시세",
    "손익계산서",
    "재무상태표",
    "현금흐름표",
    "관련뉴스",
    "출처및산식",
)

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_FILENAME_UNSAFE_RE = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')


def sanitize_excel_text(value: object) -> str:
    """외부 문자열(테마·뉴스 제목·AI 문장 등)이 엑셀 수식으로 실행되지 않도록 안전하게 정화한다."""
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    if text.startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def _safe_filename_part(value: object, fallback: str) -> str:
    cleaned = _FILENAME_UNSAFE_RE.sub("_", str(value or "").strip())
    cleaned = re.sub(r"\s+", "_", cleaned).strip("._ ")
    return cleaned[:40] if cleaned else fallback


def build_export_filename(result: AnalysisResult) -> str:
    """선택 종목명·6자리 종목코드·시세 기간·조회일이 포함된 안전한 XLSX 파일명을 생성한다."""
    stock = result.financial.stock
    theme_part = _safe_filename_part(result.theme, "테마")
    name_part = _safe_filename_part(stock.name, "종목")
    code_part = _safe_filename_part(stock.code or "000000", "000000")
    period_part = _safe_filename_part(
        result.query_options.price_period if result.query_options else "3M",
        "3M",
    )
    date_part = result.financial.retrieved_at.strftime("%Y%m%d")
    return f"주식분석_{theme_part}_{name_part}_{code_part}_{period_part}_{date_part}.xlsx"


def _to_eok_float(raw_krw: int | Decimal | float | None) -> float | None:
    if raw_krw is None or isinstance(raw_krw, bool):
        return None
    try:
        dec = Decimal(str(raw_krw)) / Decimal("100000000")
        val = float(round(dec, 2))
        return val if math.isfinite(val) else None
    except Exception:
        return None


def _to_finite_float(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        num = float(value)
        return num if math.isfinite(num) else None
    except (TypeError, ValueError):
        return None


def _write_optional_number(ws, row: int, col: int, value: object, num_format=None) -> bool:
    num = _to_finite_float(value)
    if num is None:
        ws.write_blank(row, col, None, num_format)
        return False
    ws.write_number(row, col, num, num_format)
    return True


def _resolve_roe_metric(result: AnalysisResult) -> MetricValue | None:
    if result.financial_ratios and result.financial_ratios.roe is not None:
        return result.financial_ratios.roe
    if result.financial.roe_metric is not None:
        return result.financial.roe_metric
    return None


def build_workbook_bytes(result: AnalysisResult) -> bytes:
    """AnalysisResult 단일 객체로부터 8개 시트와 셀 참조 네이티브 차트를 포함한 XLSX 바이트를 생성한다."""
    output = io.BytesIO()
    workbook = xlsxwriter.Workbook(
        output,
        {
            "in_memory": True,
            "strings_to_formulas": False,
            "strings_to_urls": False,
        },
    )

    title_fmt = workbook.add_format({"bold": True, "font_size": 13, "font_color": "#1B365D"})
    header_fmt = workbook.add_format(
        {"bold": True, "bg_color": "#E8EEF5", "border": 1, "align": "center"}
    )
    label_fmt = workbook.add_format({"bold": True, "bg_color": "#F5F7FA", "border": 1})
    text_fmt = workbook.add_format({"border": 1})
    int_fmt = workbook.add_format({"border": 1, "num_format": "#,##0"})
    dec_fmt = workbook.add_format({"border": 1, "num_format": "#,##0.00"})

    ws_summary = workbook.add_worksheet("요약")
    ws_metrics = workbook.add_worksheet("핵심지표")
    ws_history = workbook.add_worksheet("일별시세")
    ws_is = workbook.add_worksheet("손익계산서")
    ws_bs = workbook.add_worksheet("재무상태표")
    ws_cf = workbook.add_worksheet("현금흐름표")
    ws_news = workbook.add_worksheet("관련뉴스")
    ws_sources = workbook.add_worksheet("출처및산식")

    stock = result.financial.stock
    stock_code = str(stock.code or "")
    price_period = result.query_options.price_period if result.query_options else "3M"
    fs_div_pref = result.query_options.fs_div if result.query_options else "CFS"
    retrieved_str = format_timestamp(result.financial.retrieved_at)
    price_at_str = format_timestamp(result.financial.price_at, fallback="제공 시각 미확인")

    # 1. 요약 시트
    ws_summary.set_column(0, 0, 22)
    ws_summary.set_column(1, 3, 38)
    ws_summary.write_string(0, 0, "관심 테마 기반 주식 분석 요약", title_fmt)

    summary_meta = [
        ("검색 테마", sanitize_excel_text(result.theme)),
        ("선택 종목명", sanitize_excel_text(stock.name)),
        ("종목코드 (6자리)", stock_code),
        ("상장 시장", sanitize_excel_text(stock.market or "미확인")),
        ("시세 조회 범위", sanitize_excel_text(price_period)),
        ("재무제표 요청 기준", sanitize_excel_text(fs_div_pref)),
        ("조회 기준 시각", sanitize_excel_text(retrieved_str)),
        ("가격 기준 시각", sanitize_excel_text(price_at_str)),
        ("결과 식별자(result_id)", sanitize_excel_text(result.result_id)),
    ]
    for idx, (k, v) in enumerate(summary_meta, start=2):
        ws_summary.write_string(idx, 0, k, label_fmt)
        ws_summary.write_string(idx, 1, v, text_fmt)

    row_cursor = len(summary_meta) + 3
    ws_summary.write_string(row_cursor, 0, "영역별 수집 상태", title_fmt)
    row_cursor += 1
    ws_summary.write_string(row_cursor, 0, "영역", header_fmt)
    ws_summary.write_string(row_cursor, 1, "상태", header_fmt)
    for sec_key, sec_val in (result.section_status or {}).items():
        row_cursor += 1
        ws_summary.write_string(row_cursor, 0, sanitize_excel_text(sec_key), text_fmt)
        ws_summary.write_string(row_cursor, 1, sanitize_excel_text(sec_val), text_fmt)

    row_cursor += 2
    ws_summary.write_string(row_cursor, 0, "AI 종합 분석 리포트", title_fmt)
    row_cursor += 1
    ws_summary.write_string(row_cursor, 0, "구분", header_fmt)
    ws_summary.write_string(row_cursor, 1, "설명 내용", header_fmt)
    ws_summary.write_string(row_cursor, 2, "근거 (뉴스/재무 fact_id)", header_fmt)

    if result.commentary is not None:
        comm = result.commentary
        for idx, stmt in enumerate(comm.theme_trend, start=1):
            row_cursor += 1
            ws_summary.write_string(row_cursor, 0, f"테마 동향 {idx}", label_fmt)
            ws_summary.write_string(row_cursor, 1, sanitize_excel_text(stmt.text), text_fmt)
            refs = ", ".join(f"뉴스 {nid}" for nid in stmt.news_ids)
            ws_summary.write_string(row_cursor, 2, sanitize_excel_text(refs), text_fmt)

        row_cursor += 1
        ws_summary.write_string(row_cursor, 0, "선택 종목 지표 설명", label_fmt)
        ws_summary.write_string(row_cursor, 1, sanitize_excel_text(comm.stock_comment), text_fmt)
        ws_summary.write_string(row_cursor, 2, "기본 주가·PER 조회값", text_fmt)

        for idx, fstmt in enumerate(comm.financial_commentary, start=1):
            row_cursor += 1
            ws_summary.write_string(row_cursor, 0, f"기업 재무 분석 {idx}", label_fmt)
            ws_summary.write_string(row_cursor, 1, sanitize_excel_text(fstmt.text), text_fmt)
            ws_summary.write_string(
                row_cursor, 2, sanitize_excel_text(", ".join(fstmt.fact_ids)), text_fmt
            )

        for label, item in (
            ("기회 요인", comm.opportunity),
            ("위험 요인", comm.risk),
            ("종합 의견", comm.summary),
        ):
            row_cursor += 1
            ws_summary.write_string(row_cursor, 0, label, label_fmt)
            if item is not None:
                ws_summary.write_string(row_cursor, 1, sanitize_excel_text(item.text), text_fmt)
                refs = ", ".join(f"뉴스 {nid}" for nid in item.news_ids)
                ws_summary.write_string(row_cursor, 2, sanitize_excel_text(refs), text_fmt)
            else:
                ws_summary.write_string(row_cursor, 1, "관련 뉴스 근거 부족으로 미생성", text_fmt)
                ws_summary.write_string(row_cursor, 2, "-", text_fmt)
    else:
        row_cursor += 1
        ws_summary.write_string(row_cursor, 0, "AI 설명 상태", label_fmt)
        ws_summary.write_string(
            row_cursor,
            1,
            sanitize_excel_text(result.commentary_error or "AI 설명이 준비되지 않았습니다."),
            text_fmt,
        )
        ws_summary.write_string(row_cursor, 2, "-", text_fmt)

    # 2. 핵심지표 시트
    ws_metrics.set_column(0, 0, 16)
    ws_metrics.set_column(1, 2, 18)
    ws_metrics.set_column(3, 9, 22)
    ws_metrics.write_string(0, 0, f"{stock.name} ({stock_code}) 핵심 금융·재무 지표", title_fmt)
    ws_metrics.write_string(1, 0, "종목코드", label_fmt)
    ws_metrics.write_string(1, 1, stock_code, text_fmt)

    metric_headers = (
        "지표명",
        "수치 원값",
        "화면 표시값",
        "단위",
        "통화",
        "기준 기간/시각",
        "제공/계산 구분",
        "산식/원필드",
        "출처",
        "상태/사유",
    )
    for col_idx, h in enumerate(metric_headers):
        ws_metrics.write_string(3, col_idx, h, header_fmt)

    fin = result.financial
    roe_metric = _resolve_roe_metric(result)
    ratios = result.financial_ratios
    op_margin = ratios.operating_margin if ratios else None
    net_margin = ratios.net_margin if ratios else None
    debt_ratio = ratios.debt_ratio if ratios else None

    metric_rows = [
        (
            "현재 주가",
            fin.price,
            format_price(fin.price),
            "원",
            "KRW",
            price_at_str,
            "provided",
            "regularMarketPrice / close",
            fin.source,
            "ok" if fin.price is not None else "가격 정보 없음",
        ),
        (
            "PER",
            round(fin.per, 2) if fin.per is not None else None,
            format_multiple(fin.per),
            "배",
            "",
            fin.per_period_end or "TTM",
            "calculated" if fin.per_eps is not None else "provided",
            "현재 주가 / 최근 4개 분기 합산 EPS" if fin.per_eps is not None else "trailingPE",
            fin.source,
            "ok" if fin.per is not None else "정보 없음",
        ),
        (
            "PBR",
            round(fin.pbr, 2) if fin.pbr is not None else None,
            format_multiple(fin.pbr),
            "배",
            "",
            fin.pbr_metric.period if fin.pbr_metric else "-",
            fin.pbr_metric.basis if fin.pbr_metric else "missing",
            (fin.pbr_metric.formula or fin.pbr_metric.raw_field) if fin.pbr_metric else "-",
            fin.pbr_metric.source if fin.pbr_metric else fin.source,
            (fin.pbr_metric.message or fin.pbr_metric.status) if fin.pbr_metric else "정보 없음",
        ),
        (
            "ROE",
            round(roe_metric.value, 2) if (roe_metric and roe_metric.value is not None) else None,
            format_percent(roe_metric.value if roe_metric else None),
            "%",
            "",
            roe_metric.period if roe_metric else "-",
            roe_metric.basis if roe_metric else "missing",
            (roe_metric.formula or roe_metric.raw_field) if roe_metric else "-",
            roe_metric.source if roe_metric else "OpenDART",
            (roe_metric.message or roe_metric.status) if roe_metric else "정보 없음",
        ),
        (
            "EPS",
            round(fin.eps, 2) if fin.eps is not None else None,
            format_per_share_krw(fin.eps),
            "원",
            "KRW",
            fin.eps_metric.period if fin.eps_metric else "-",
            fin.eps_metric.basis if fin.eps_metric else "missing",
            fin.eps_metric.raw_field if fin.eps_metric else "-",
            fin.eps_metric.source if fin.eps_metric else fin.source,
            (fin.eps_metric.message or fin.eps_metric.status) if fin.eps_metric else "정보 없음",
        ),
        (
            "BPS",
            round(fin.bps, 2) if fin.bps is not None else None,
            format_per_share_krw(fin.bps),
            "원",
            "KRW",
            fin.bps_metric.period if fin.bps_metric else "-",
            fin.bps_metric.basis if fin.bps_metric else "missing",
            fin.bps_metric.raw_field if fin.bps_metric else "-",
            fin.bps_metric.source if fin.bps_metric else fin.source,
            (fin.bps_metric.message or fin.bps_metric.status) if fin.bps_metric else "정보 없음",
        ),
        (
            "영업이익률",
            round(op_margin.value, 2) if (op_margin and op_margin.value is not None) else None,
            format_percent(op_margin.value if op_margin else None),
            "%",
            "",
            op_margin.period if op_margin else "-",
            op_margin.basis if op_margin else "missing",
            op_margin.formula if op_margin else "영업이익 / 매출액 × 100",
            op_margin.source if op_margin else "OpenDART",
            (op_margin.message or op_margin.status) if op_margin else "정보 없음",
        ),
        (
            "순이익률",
            round(net_margin.value, 2) if (net_margin and net_margin.value is not None) else None,
            format_percent(net_margin.value if net_margin else None),
            "%",
            "",
            net_margin.period if net_margin else "-",
            net_margin.basis if net_margin else "missing",
            net_margin.formula if net_margin else "당기순이익 / 매출액 × 100",
            net_margin.source if net_margin else "OpenDART",
            (net_margin.message or net_margin.status) if net_margin else "정보 없음",
        ),
        (
            "부채비율",
            round(debt_ratio.value, 2) if (debt_ratio and debt_ratio.value is not None) else None,
            format_percent(debt_ratio.value if debt_ratio else None),
            "%",
            "",
            debt_ratio.period if debt_ratio else "-",
            debt_ratio.basis if debt_ratio else "missing",
            debt_ratio.formula if debt_ratio else "부채총계 / 자본총계 × 100",
            debt_ratio.source if debt_ratio else "OpenDART",
            (debt_ratio.message or debt_ratio.status) if debt_ratio else "정보 없음",
        ),
    ]

    for r_offset, m_row in enumerate(metric_rows, start=4):
        ws_metrics.write_string(r_offset, 0, sanitize_excel_text(m_row[0]), label_fmt)
        _write_optional_number(ws_metrics, r_offset, 1, m_row[1], dec_fmt)
        for c_idx in range(2, 10):
            ws_metrics.write_string(r_offset, c_idx, sanitize_excel_text(m_row[c_idx]), text_fmt)

    # 3. 일별시세 시트 + 셀 참조 주가/거래량 네이티브 차트
    ws_history.set_column(0, 0, 14)
    ws_history.set_column(1, 5, 16)
    ws_history.set_column(6, 8, 22)
    ws_history.write_string(0, 0, f"{stock.name} ({stock_code}) 일별 주가·거래량 시세", title_fmt)

    hist = result.price_history
    hist_meta = (
        f"기간: {hist.period if hist else price_period} · "
        f"통화: {hist.currency if hist else 'KRW'} · "
        f"조정 기준: {'미수정 원시 시세(auto_adjust=False)' if (hist and not hist.adjusted) else '수정주가'} · "
        f"출처: {hist.source if hist else fin.source}"
    )
    ws_history.write_string(1, 0, sanitize_excel_text(hist_meta))

    hist_headers = (
        "거래일",
        "시가(원)",
        "고가(원)",
        "저가(원)",
        "종가(원)",
        "거래량(주)",
        "통화",
        "조정여부",
        "출처",
    )
    for col_idx, h in enumerate(hist_headers):
        ws_history.write_string(3, col_idx, h, header_fmt)

    bars = hist.bars if (hist and hist.status == "ok" and hist.bars) else ()
    if bars:
        for idx, bar in enumerate(bars, start=4):
            ws_history.write_string(idx, 0, sanitize_excel_text(bar.trade_date), text_fmt)
            _write_optional_number(ws_history, idx, 1, bar.open, int_fmt)
            _write_optional_number(ws_history, idx, 2, bar.high, int_fmt)
            _write_optional_number(ws_history, idx, 3, bar.low, int_fmt)
            _write_optional_number(ws_history, idx, 4, bar.close, int_fmt)
            _write_optional_number(ws_history, idx, 5, bar.volume, int_fmt)
            ws_history.write_string(idx, 6, sanitize_excel_text(hist.currency), text_fmt)
            ws_history.write_string(
                idx, 7, "미수정(raw)" if not hist.adjusted else "수정", text_fmt
            )
            ws_history.write_string(idx, 8, sanitize_excel_text(hist.source), text_fmt)

        last_hist_row = 3 + len(bars)  # 0-based index of last data row
        # 네이티브 선형 주가 차트 (일별시세 시트 셀 참조)
        price_chart = workbook.add_chart({"type": "line"})
        price_chart.add_series(
            {
                "name": "=일별시세!$E$4",
                "categories": ["일별시세", 4, 0, last_hist_row, 0],
                "values": ["일별시세", 4, 4, last_hist_row, 4],
                "line": {"color": "#1F77B4", "width": 2.0},
            }
        )
        price_chart.set_title({"name": f"{stock.name} ({stock_code}) 일별 종가 추이 (원)"})
        price_chart.set_x_axis({"name": "거래일"})
        price_chart.set_y_axis({"name": "종가 (원)"})
        price_chart.set_size({"width": 640, "height": 300})
        ws_history.insert_chart("J4", price_chart)

        # 네이티브 막대 거래량 차트 (일별시세 시트 셀 참조)
        vol_chart = workbook.add_chart({"type": "column"})
        vol_chart.add_series(
            {
                "name": "=일별시세!$F$4",
                "categories": ["일별시세", 4, 0, last_hist_row, 0],
                "values": ["일별시세", 4, 5, last_hist_row, 5],
                "fill": {"color": "#7F8C8D"},
            }
        )
        vol_chart.set_title({"name": f"{stock.name} ({stock_code}) 일별 거래량 추이 (주)"})
        vol_chart.set_x_axis({"name": "거래일"})
        vol_chart.set_y_axis({"name": "거래량 (주)"})
        vol_chart.set_size({"width": 640, "height": 260})
        ws_history.insert_chart("J20", vol_chart)
    else:
        empty_reason = (hist.message if hist and hist.message else "수집된 일별 시세 이력이 없습니다.")
        ws_history.write_string(4, 0, "자료 없음", label_fmt)
        ws_history.write_string(4, 1, sanitize_excel_text(empty_reason), text_fmt)

    # 4. 손익계산서 시트 + 셀 참조 연간 실적 네이티브 차트
    ws_is.set_column(0, 1, 14)
    ws_is.set_column(2, 7, 20)
    ws_is.set_column(8, 10, 22)
    ws_is.write_string(0, 0, f"{stock.name} ({stock_code}) 손익계산서 (연간 비교 및 상세)", title_fmt)

    is_headers = (
        "사업연도",
        "재무구분",
        "매출액(원)",
        "영업이익(원)",
        "당기순이익(원)",
        "매출액(억원)",
        "영업이익(억원)",
        "당기순이익(억원)",
        "공시접수번호",
        "출처",
        "비고",
    )
    for col_idx, h in enumerate(is_headers):
        ws_is.write_string(3, col_idx, h, header_fmt)

    annuals = result.annual_summaries or ()
    if annuals:
        for idx, ann in enumerate(annuals, start=4):
            ws_is.write_string(idx, 0, sanitize_excel_text(f"{ann.bsns_year}년"), text_fmt)
            ws_is.write_string(idx, 1, sanitize_excel_text(ann.fs_div), text_fmt)
            _write_optional_number(ws_is, idx, 2, ann.revenue, int_fmt)
            _write_optional_number(ws_is, idx, 3, ann.operating_income, int_fmt)
            _write_optional_number(ws_is, idx, 4, ann.net_income, int_fmt)
            _write_optional_number(ws_is, idx, 5, _to_eok_float(ann.revenue), dec_fmt)
            _write_optional_number(ws_is, idx, 6, _to_eok_float(ann.operating_income), dec_fmt)
            _write_optional_number(ws_is, idx, 7, _to_eok_float(ann.net_income), dec_fmt)
            ws_is.write_string(idx, 8, sanitize_excel_text(ann.rcept_no or "-"), text_fmt)
            ws_is.write_string(idx, 9, "OpenDART", text_fmt)
            note = "별도(OFS) 대체" if ann.is_separate_fallback else "정상"
            ws_is.write_string(idx, 10, note, text_fmt)

        last_is_row = 3 + len(annuals)
        income_chart = workbook.add_chart({"type": "column"})
        for col_offset, color in ((5, "#2B5C8F"), (6, "#3690C0"), (7, "#67A9CF")):
            income_chart.add_series(
                {
                    "name": ["손익계산서", 3, col_offset],
                    "categories": ["손익계산서", 4, 0, last_is_row, 0],
                    "values": ["손익계산서", 4, col_offset, last_is_row, col_offset],
                    "fill": {"color": color},
                }
            )
        income_chart.set_title(
            {"name": f"{stock.name} ({stock_code}) 연도별 실적 추이 (억원, {annuals[-1].fs_div})"}
        )
        income_chart.set_x_axis({"name": "사업연도"})
        income_chart.set_y_axis({"name": "금액 (억원)"})
        income_chart.set_size({"width": 620, "height": 300})
        ws_is.insert_chart("L4", income_chart)

        # 요약 시트에도 동일한 셀 참조 실적 차트를 제공
        summary_income_chart = workbook.add_chart({"type": "column"})
        for col_offset, color in ((5, "#2B5C8F"), (6, "#3690C0"), (7, "#67A9CF")):
            summary_income_chart.add_series(
                {
                    "name": ["손익계산서", 3, col_offset],
                    "categories": ["손익계산서", 4, 0, last_is_row, 0],
                    "values": ["손익계산서", 4, col_offset, last_is_row, col_offset],
                    "fill": {"color": color},
                }
            )
        summary_income_chart.set_title(
            {"name": f"{stock.name} ({stock_code}) 연도별 매출·영업이익·순이익 (억원)"}
        )
        summary_income_chart.set_size({"width": 580, "height": 280})
        ws_summary.insert_chart("E2", summary_income_chart)
    else:
        ws_is.write_string(4, 0, "자료 없음", label_fmt)
        ws_is.write_string(
            4, 1, "OpenDART 연간 손익계산서 자료가 없거나 미공시 상태입니다.", text_fmt
        )

    # 5. 재무상태표 시트 + 셀 참조 재무구조 차트
    ws_bs.set_column(0, 1, 14)
    ws_bs.set_column(2, 7, 20)
    ws_bs.set_column(8, 10, 22)
    ws_bs.write_string(0, 0, f"{stock.name} ({stock_code}) 재무상태표 (연도별 기말 잔액)", title_fmt)

    bs_headers = (
        "사업연도",
        "재무구분",
        "자산총계(원)",
        "부채총계(원)",
        "자본총계(원)",
        "자산총계(억원)",
        "부채총계(억원)",
        "자본총계(억원)",
        "공시접수번호",
        "출처",
        "비고",
    )
    for col_idx, h in enumerate(bs_headers):
        ws_bs.write_string(3, col_idx, h, header_fmt)

    if annuals:
        for idx, ann in enumerate(annuals, start=4):
            ws_bs.write_string(idx, 0, sanitize_excel_text(f"{ann.bsns_year}년"), text_fmt)
            ws_bs.write_string(idx, 1, sanitize_excel_text(ann.fs_div), text_fmt)
            _write_optional_number(ws_bs, idx, 2, ann.total_assets, int_fmt)
            _write_optional_number(ws_bs, idx, 3, ann.total_liabilities, int_fmt)
            _write_optional_number(ws_bs, idx, 4, ann.total_equity, int_fmt)
            _write_optional_number(ws_bs, idx, 5, _to_eok_float(ann.total_assets), dec_fmt)
            _write_optional_number(ws_bs, idx, 6, _to_eok_float(ann.total_liabilities), dec_fmt)
            _write_optional_number(ws_bs, idx, 7, _to_eok_float(ann.total_equity), dec_fmt)
            ws_bs.write_string(idx, 8, sanitize_excel_text(ann.rcept_no or "-"), text_fmt)
            ws_bs.write_string(idx, 9, "OpenDART", text_fmt)
            eq_note = (
                "자본잠식(음수 자본)"
                if (ann.total_equity is not None and ann.total_equity < 0)
                else "기말 잔액"
            )
            ws_bs.write_string(idx, 10, eq_note, text_fmt)

        last_bs_row = 3 + len(annuals)
        bs_chart = workbook.add_chart({"type": "column"})
        for col_offset, color in ((5, "#34495E"), (6, "#E67E22"), (7, "#27AE60")):
            bs_chart.add_series(
                {
                    "name": ["재무상태표", 3, col_offset],
                    "categories": ["재무상태표", 4, 0, last_bs_row, 0],
                    "values": ["재무상태표", 4, col_offset, last_bs_row, col_offset],
                    "fill": {"color": color},
                }
            )
        bs_chart.set_title(
            {"name": f"{stock.name} ({stock_code}) 연도별 자산·부채·자본 (억원, {annuals[-1].fs_div})"}
        )
        bs_chart.set_x_axis({"name": "사업연도"})
        bs_chart.set_y_axis({"name": "금액 (억원)"})
        bs_chart.set_size({"width": 620, "height": 300})
        ws_bs.insert_chart("L4", bs_chart)
    else:
        ws_bs.write_string(4, 0, "자료 없음", label_fmt)
        ws_bs.write_string(
            4, 1, "OpenDART 재무상태표 자료가 없거나 미공시 상태입니다.", text_fmt
        )

    # 6. 현금흐름표 시트
    ws_cf.set_column(0, 1, 14)
    ws_cf.set_column(2, 7, 22)
    ws_cf.set_column(8, 10, 22)
    ws_cf.write_string(0, 0, f"{stock.name} ({stock_code}) 현금흐름표 (연도별 비교)", title_fmt)

    cf_headers = (
        "사업연도",
        "재무구분",
        "영업활동현금흐름(원)",
        "투자활동현금흐름(원)",
        "재무활동현금흐름(원)",
        "영업활동현금흐름(억원)",
        "투자활동현금흐름(억원)",
        "재무활동현금흐름(억원)",
        "공시접수번호",
        "출처",
        "해석 유의사항",
    )
    for col_idx, h in enumerate(cf_headers):
        ws_cf.write_string(3, col_idx, h, header_fmt)

    if annuals:
        for idx, ann in enumerate(annuals, start=4):
            ws_cf.write_string(idx, 0, sanitize_excel_text(f"{ann.bsns_year}년"), text_fmt)
            ws_cf.write_string(idx, 1, sanitize_excel_text(ann.fs_div), text_fmt)
            _write_optional_number(ws_cf, idx, 2, ann.operating_cash_flow, int_fmt)
            _write_optional_number(ws_cf, idx, 3, ann.investing_cash_flow, int_fmt)
            _write_optional_number(ws_cf, idx, 4, ann.financing_cash_flow, int_fmt)
            _write_optional_number(ws_cf, idx, 5, _to_eok_float(ann.operating_cash_flow), dec_fmt)
            _write_optional_number(ws_cf, idx, 6, _to_eok_float(ann.investing_cash_flow), dec_fmt)
            _write_optional_number(ws_cf, idx, 7, _to_eok_float(ann.financing_cash_flow), dec_fmt)
            ws_cf.write_string(idx, 8, sanitize_excel_text(ann.rcept_no or "-"), text_fmt)
            ws_cf.write_string(idx, 9, "OpenDART", text_fmt)
            ws_cf.write_string(
                idx,
                10,
                "투자활동 현금흐름 음수(-)는 설비·연구 투자 유출을 의미할 수 있으며 적자와 다릅니다.",
                text_fmt,
            )
    else:
        ws_cf.write_string(4, 0, "자료 없음", label_fmt)
        ws_cf.write_string(
            4, 1, "OpenDART 현금흐름표 자료가 없거나 미공시 상태입니다.", text_fmt
        )

    # 7. 관련뉴스 시트
    ws_news.set_column(0, 0, 10)
    ws_news.set_column(1, 1, 55)
    ws_news.set_column(2, 3, 22)
    ws_news.set_column(4, 4, 48)
    ws_news.write_string(0, 0, f"{stock.name} ({stock_code}) 관련 뉴스 목록", title_fmt)
    ws_news.write_string(
        1,
        0,
        sanitize_excel_text(
            f"출처: {result.news.source} · 조회 시각: {format_timestamp(result.news.retrieved_at)}"
        ),
    )

    news_headers = ("뉴스 번호", "기사 제목", "발행 시각", "언론사 도메인", "안전 확인된 원문 URL")
    for col_idx, h in enumerate(news_headers):
        ws_news.write_string(3, col_idx, h, header_fmt)

    if result.news.articles:
        for idx, art in enumerate(result.news.articles, start=4):
            ws_news.write_number(idx, 0, art.id, int_fmt)
            ws_news.write_string(idx, 1, sanitize_excel_text(art.title), text_fmt)
            ws_news.write_string(
                idx, 2, sanitize_excel_text(art.published_at or "미확인"), text_fmt
            )
            ws_news.write_string(
                idx, 3, sanitize_excel_text(art.publisher_domain or "미확인"), text_fmt
            )
            safe_url = safe_news_url(art.url) or "미확인"
            ws_news.write_string(idx, 4, sanitize_excel_text(safe_url), text_fmt)
    else:
        ws_news.write_string(4, 0, "자료 없음", label_fmt)
        ws_news.write_string(
            4, 1, sanitize_excel_text(result.news.message or "표시할 관련 뉴스가 없습니다."), text_fmt
        )

    # 8. 출처및산식 시트
    ws_sources.set_column(0, 1, 24)
    ws_sources.set_column(2, 4, 38)
    ws_sources.write_string(0, 0, "데이터 출처·기준 시각·산식 안내", title_fmt)

    source_headers = ("데이터 영역", "제공 출처", "기준 기간/시각", "계산 산식 및 기준", "주의사항")
    for col_idx, h in enumerate(source_headers):
        ws_sources.write_string(2, col_idx, h, header_fmt)

    source_rows = [
        (
            "상장 종목 대조",
            "KRX 상장법인 목록 (FinanceDataReader)",
            retrieved_str,
            "공식 종목명 정확 일치 → 6자리 종목코드 및 시장(.KS/.KQ) 매핑",
            "상장 미확인 후보나 지원 대상 외 시장은 상세 조회 전 거부합니다.",
        ),
        (
            "기본 주가·PER·EPS·BPS·PBR",
            fin.source,
            f"조회: {retrieved_str} / 가격: {price_at_str}",
            "제공 원필드 우선 사용. PBR 대체 계산 시: 현재 주가 / BPS (동일 통화·양수 BPS 검증 시)",
            "결측값(None)은 0으로 대체하지 않고 빈 셀과 사유로 표시합니다.",
        ),
        (
            "일별 주가·거래량 시세",
            hist.source if hist else fin.source,
            f"기간: {price_period}",
            "일별 시가·고가·저가·종가·거래량 오름차순 (auto_adjust=False 원시 시세 기준)",
            "휴장일을 임의 보간하지 않으며 관련 종목 등락은 동일 조정 기준 직전 거래일 대비로 계산합니다.",
        ),
        (
            "기업 재무제표 (IS/CIS/BS/CF)",
            "금융감독원 전자공시시스템 (OpenDART fnlttSinglAcntAll)",
            f"최근 공시 사업보고서 ({fs_div_pref} 우선)",
            "원 단위 정수 정밀도 보존, 억원 표시는 원값 / 100,000,000 (소수 둘째 자리 반올림)",
            "연결(CFS)과 별도(OFS)를 혼합하지 않으며, BS 기말 잔액은 분기 차감하지 않습니다.",
        ),
        (
            "재무 비율 (ROE·이익률·부채비율)",
            "OpenDART 주요재무지표(fnlttSinglIndx) 및 재무제표 계산",
            ratios.bsns_year + "년" if (ratios and ratios.bsns_year) else "최근 사업보고서",
            "영업이익률=영업이익/매출액×100, 순이익률=당기순이익/매출액×100, 부채비율=부채총계/자본총계×100",
            "매출 0/누락 또는 자본 0/음수(자본잠식)인 경우 비율을 임의 계산하지 않습니다.",
        ),
        (
            "관련 뉴스 및 AI 분석",
            f"{result.news.source} / Groq LLM",
            format_timestamp(result.news.retrieved_at),
            "뉴스 제목(news_ids) 및 재무 사실(fact_id) 인용 검증",
            "본 서비스는 정보 탐색 보조 도구이며 매수·매도 등 투자 권유나 수익을 보장하지 않습니다.",
        ),
    ]
    for r_idx, s_row in enumerate(source_rows, start=3):
        ws_sources.write_string(r_idx, 0, s_row[0], label_fmt)
        for c_idx in range(1, 5):
            ws_sources.write_string(r_idx, c_idx, sanitize_excel_text(s_row[c_idx]), text_fmt)

    workbook.close()
    return output.getvalue()


def reconcile_workbook_values(result: AnalysisResult, xlsx_bytes: bytes) -> dict[str, object]:
    """생성된 XLSX 바이트를 열어 시트 구성·종목코드 선행0·수치 원자료·네이티브 차트 셀 참조를 검증한다."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    sheet_names = tuple(wb.sheetnames)

    ws_summary = wb["요약"]
    summary_code = ws_summary["B5"].value
    summary_code_type = ws_summary["B5"].data_type

    ws_metrics = wb["핵심지표"]
    metrics_code = ws_metrics["B2"].value
    metric_values: dict[str, object] = {}
    for row in range(5, 14):
        name = ws_metrics.cell(row=row, column=1).value
        val = ws_metrics.cell(row=row, column=2).value
        if name:
            metric_values[str(name)] = val

    ws_history = wb["일별시세"]
    history_rows = 0
    first_close = None
    last_close = None
    if result.price_history and result.price_history.bars:
        history_rows = len(result.price_history.bars)
        first_close = ws_history.cell(row=5, column=5).value
        last_close = ws_history.cell(row=4 + history_rows, column=5).value

    ws_is = wb["손익계산서"]
    annual_rows = len(result.annual_summaries or ())
    latest_rev_krw = None
    latest_rev_eok = None
    if annual_rows > 0:
        latest_rev_krw = ws_is.cell(row=4 + annual_rows, column=3).value
        latest_rev_eok = ws_is.cell(row=4 + annual_rows, column=6).value

    # 수식 주입 여부 확인: 모든 셀에 의도치 않은 수식(data_type == 'f')이 없어야 함
    formula_cells: list[str] = []
    for sname in wb.sheetnames:
        ws = wb[sname]
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_cells.append(f"{sname}!{cell.coordinate}")

    # OOXML 차트 XML 내부의 셀 참조(c:f) 확인
    chart_references: list[str] = []
    with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as zf:
        chart_files = [n for n in zf.namelist() if n.startswith("xl/charts/chart") and n.endswith(".xml")]
        for cf in sorted(chart_files):
            root = ET.fromstring(zf.read(cf))
            for elem in root.iter():
                if elem.tag.endswith("}f") and elem.text:
                    chart_references.append(elem.text.strip())

    expected_code = str(result.financial.stock.code or "")
    code_matches = (
        summary_code == expected_code
        and metrics_code == expected_code
        and summary_code_type == "s"
    )
    price_matches = metric_values.get("현재 주가") == result.financial.price
    expected_per = round(result.financial.per, 2) if result.financial.per is not None else None
    per_matches = metric_values.get("PER") == expected_per

    return {
        "sheet_names": sheet_names,
        "all_sheets_present": sheet_names == SHEET_NAMES,
        "code_matches": code_matches,
        "summary_code": summary_code,
        "price_matches": price_matches,
        "per_matches": per_matches,
        "metric_values": metric_values,
        "history_rows": history_rows,
        "first_close": first_close,
        "last_close": last_close,
        "annual_rows": annual_rows,
        "latest_rev_krw": latest_rev_krw,
        "latest_rev_eok": latest_rev_eok,
        "formula_cells": formula_cells,
        "chart_count": len(chart_files),
        "chart_references": chart_references,
    }
