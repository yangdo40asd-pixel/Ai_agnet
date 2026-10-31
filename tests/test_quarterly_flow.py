"""분기 요청의 실제 pipeline/UI 연결, 기간 문맥, 연간·캔버스 보존 회귀."""
from dataclasses import replace
from datetime import timedelta
import unittest
from unittest.mock import patch

from src.analysis_flow import needs_data
from src.config import Settings
from src.models import now_kst
from src.data.financials import build_quarter_comparison
from src.pipeline import analyze_selected_stock
from src.agents.chat import answer_question
from src.agents.answer_quality import QUALITY_CHECKS
from scripts.llm_fixture import synthetic_result
from test_quarterly import statement, evidence


def bundle(year=2025):
    statements = (statement(str(year-1), revenue="10000000000", profit="1000000000"), statement(str(year)))
    return {"status": "ok", "statements": statements, "summaries": build_quarter_comparison(statements), "calls": ()}


def fake_llm(settings, instruction, payload, *args, **kwargs):
    if payload["task"] == "chat_answer":
        from src.agents.financial_numbers import verified_sentence
        refs = {k: v for k, v in payload["evidence"].items() if v.get("kind") == "quarterly_comparison"}
        return {"text": "수익성이 개선된 실적입니다. " + " ".join(verified_sentence(v) for v in refs.values()),
                "evidence_ids": list(refs), "follow_up": "", "blocks": []}
    if payload["task"] == "chat_grounding_review":
        return {"checks": dict.fromkeys(QUALITY_CHECKS, True), "supported": True, "answers_question": True,
                "unit_reviews": [{"id": u["id"], "supported": True, "on_topic": True,
                                  "relevance_evidence_ids": u["evidence_ids"], "reason": "가상 검사"}
                                 for u in payload["review_units"]]}
    raise AssertionError(payload["task"])


class QuarterlyFlowTests(unittest.TestCase):
    def setUp(self):
        self.settings = Settings("", "", "", "", dart_api_key="fixture-key", news_relevance_mode="strict")
        self.old = synthetic_result()

    def test_pipeline_adds_quarters_without_overwriting_annual_snapshot(self):
        with patch("src.data.quarterly.collect_quarterly_financials", return_value=bundle()) as fetch, \
             patch("src.pipeline.collect_company_financials") as annual:
            result = analyze_selected_stock("반도체", self.old.candidates, self.old.target_code,
                settings=self.settings, scope="financials", previous=self.old, question="2025년 3분기 실적은?")
        fetch.assert_called_once(); annual.assert_not_called()
        self.assertEqual(result.annual_summaries, self.old.annual_summaries)
        self.assertTrue(result.quarterly_summaries)
        self.assertNotEqual(result.result_id, self.old.result_id)
        self.assertFalse(self.old.quarterly_summaries)
        self.assertEqual(result.section_status["quarterly_request"], "2025Q3:auto:3m:yoy")

    def test_quarter_plus_price_collects_requested_sections_without_basic_news_report(self):
        with patch("src.data.quarterly.collect_quarterly_financials", return_value=bundle()), \
             patch("src.pipeline.fetch_financial", return_value=self.old.financial) as finance, \
             patch("src.pipeline.fetch_news") as news, patch("src.pipeline.fetch_price_history") as history, \
             patch("src.pipeline.generate_commentary") as report:
            result = analyze_selected_stock("반도체", self.old.candidates, self.old.target_code,
                settings=self.settings, scope="all", previous=self.old,
                question="2025년 3분기 실적과 현재 주가, PER을 함께 설명해줘")
        finance.assert_called_once(); news.assert_not_called(); history.assert_not_called(); report.assert_not_called()
        self.assertTrue(result.quarterly_summaries)
        self.assertEqual(result.section_status["commentary"], "not_requested")

    def test_same_period_reuses_data_and_new_or_expired_period_collects(self):
        result = replace(self.old, section_status={"dart": "ok", "quarterly_request": "2025Q3:auto:3m:yoy",
                                                 "quarterly_checked_at": now_kst().isoformat()})
        self.assertFalse(needs_data(result, "financials", "2025년 3분기 매출", now=now_kst(), dart_enabled=True))
        self.assertTrue(needs_data(result, "financials", "2024년 3분기 매출", now=now_kst(), dart_enabled=True))
        self.assertTrue(needs_data(result, "financials", "2025년 3분기 매출", now=now_kst()+timedelta(minutes=16), dart_enabled=True))
        messages = [{"role": "user", "text": "2025년 3분기 실적은?"}]
        self.assertFalse(needs_data(result, "financials", "그럼 영업이익률은?", now=now_kst(), dart_enabled=True, messages=messages))

    def test_answer_uses_quarterly_data_and_still_runs_semantic_review(self):
        result = replace(self.old, quarterly_summaries=bundle()["summaries"], financial_statements=bundle()["statements"])
        with patch("src.agents.chat._invoke", side_effect=fake_llm) as invoke:
            response = answer_question("2025년 3분기 실적은?", result, [], self.settings)
        self.assertIn("수익성이 개선", response.text)
        self.assertEqual(invoke.call_count, 2)
        self.assertTrue(all(e["kind"] == "quarterly_comparison" for e in response.evidence))
        self.assertTrue(invoke.call_args.args[2]["displayed_data_sources"])

    def test_followup_does_not_switch_back_to_annual_or_carry_into_named_company(self):
        from src.agents.quarterly import contextual_quarter
        history = [{"role": "user", "text": "2025년 3분기 실적은?"}]
        self.assertIsNone(contextual_quarter("현대차 영업이익률은?", history))
        result = replace(self.old, quarterly_summaries=bundle()["summaries"])
        with patch("src.agents.chat._invoke", side_effect=fake_llm):
            response = answer_question("그럼 영업이익률은?", result, history, self.settings)
        self.assertTrue(all(e["quarter"] == 3 for e in response.evidence))
        self.assertIn("10.00%p", response.text)

    def test_missing_current_period_remains_explicit_without_llm(self):
        result = replace(self.old, quarterly_summaries=bundle(2024)["summaries"])
        with patch("src.agents.chat._invoke") as invoke:
            response = answer_question("2025년 3분기 실적은?", result, [], self.settings)
        self.assertIn("확인하지 못했습니다", response[0]); invoke.assert_not_called()

    def test_excel_keeps_annual_rows_and_adds_actual_quarter_values_separately(self):
        import io
        import openpyxl
        from src.export.excel import build_workbook_bytes
        result = replace(self.old, quarterly_summaries=bundle()["summaries"])
        workbook = openpyxl.load_workbook(io.BytesIO(build_workbook_bytes(result)), data_only=True)
        sheet = workbook["손익계산서"]
        self.assertEqual(sheet["A5"].value, "2024년")
        self.assertEqual(sheet["C5"].value, self.old.annual_summaries[0].revenue)
        rows = {row[0]: row for row in sheet.iter_rows(values_only=True) if row[0] in {"2024 3Q", "2025 3Q"}}
        self.assertEqual(rows["2025 3Q"][2], 12000000000)
        self.assertEqual(rows["2025 3Q"][5], 120)
        self.assertIn("3개월", rows["2025 3Q"][10])
        self.assertEqual(len(workbook.sheetnames), 8)

    def test_failed_collection_is_not_presented_as_missing_report(self):
        previous = replace(self.old, section_status={"dart_quarterly": "failed", "quarterly_request": "2025Q3:auto:3m:yoy",
                                                   "quarterly_message": "외부 조회 시간 초과"})
        with patch("src.agents.chat._invoke") as invoke:
            response = answer_question("2025년 3분기 실적은?", previous, [], self.settings)
        self.assertIn("조회에 실패", response[0]); invoke.assert_not_called()

    def test_failed_refresh_removes_old_requested_quarter_from_new_snapshot(self):
        previous = replace(self.old, quarterly_summaries=bundle()["summaries"])
        failed = {"status": "failed", "statements": (), "summaries": (), "calls": ()}
        with patch("src.data.quarterly.collect_quarterly_financials", return_value=failed):
            result = analyze_selected_stock("반도체", previous.candidates, previous.target_code,
                settings=self.settings, scope="financials", previous=previous, question="2025년 3분기 실적 재조회")
        self.assertFalse(result.quarterly_summaries)
        self.assertTrue(previous.quarterly_summaries)
        self.assertEqual(result.section_status["dart_quarterly"], "failed")


class QuarterlyUiTests(unittest.TestCase):
    def setUp(self):
        import test_user_flow
        self.flow = test_user_flow.UserFlowTests("runTest")
        self.flow.setUp(); self.addCleanup(self.flow.doCleanups)
        self.settings = Settings("", "", "", "", dart_api_key="fixture-key", news_relevance_mode="strict")
        self.settings_patch = patch("src.config.load_settings", return_value=self.settings)
        self.settings_patch.start(); self.addCleanup(self.settings_patch.stop)
        self.pipeline_settings = patch("src.pipeline.load_settings", return_value=self.settings)
        self.pipeline_settings.start(); self.addCleanup(self.pipeline_settings.stop)
        self.collect_patch = patch("src.data.quarterly.collect_quarterly_financials",
            side_effect=lambda stock, request, settings: bundle(request.year))
        self.collect = self.collect_patch.start(); self.addCleanup(self.collect_patch.stop)
        self.llm_patch = patch("src.agents.chat._invoke", side_effect=fake_llm)
        self.llm_patch.start(); self.addCleanup(self.llm_patch.stop)

    def test_quarter_query_and_followup_reuse_data_and_preserve_old_canvas(self):
        self.flow.submit_theme("예시회사 2025년 3분기 실적은?")
        state = self.flow.app.session_state
        first = state["dashboard_result"]
        self.assertTrue(first.quarterly_summaries)
        self.assertEqual(state["chat_request"].original_question, "예시회사 2025년 3분기 실적은?")
        self.flow.open_canvas()
        card = state["canvases"][0]
        self.flow.submit_theme("그럼 영업이익률은?")
        self.assertEqual(self.collect.call_count, 1)
        self.assertEqual(state["dashboard_result"].result_id, first.result_id)
        self.assertEqual(state["canvases"][0]["result_id"], card["result_id"])
        self.assertEqual(state["chat_request"].status, "completed")
        self.assertTrue(all(e["kind"].startswith("quarterly_") for e in state["chat_messages"][-1]["evidence"]))
        self.flow.submit_theme("2024년 3분기 실적은?")
        self.assertEqual(self.collect.call_count, 2)
        self.assertNotEqual(state["dashboard_result"].result_id, first.result_id)
        self.assertEqual(state["canvases"][0]["result"].result_id, first.result_id)

    def test_rerun_does_not_duplicate_answer_or_recollect(self):
        self.flow.submit_theme("예시회사 2025년 3분기 실적은?")
        state = self.flow.app.session_state
        count = len(state["chat_messages"])
        self.flow.app.run()
        self.assertEqual(self.collect.call_count, 1)
        self.assertEqual(len(state["chat_messages"]), count)
